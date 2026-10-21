"""CEUS submission package: cover letter, highlights, declarations,
data/code availability, supplement, inline docx, readme, zip.
Frozen values only."""
import csv, os, zipfile, glob, shutil
from docx import Document
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
mv = {r["key"]: r["value"] for r in
      csv.DictReader(open(f"{ROOT}/analysis/manuscript_values_v3.csv"))}

def fmt(v, nd=2):
    try:
        return f"{float(v):.{nd}f}"
    except Exception:
        return str(v)

# ---------------- highlights (Elsevier: <=85 chars, 3-5) ----------------
hl = [
    "Hidden low-voltage networks reconstructed from open data in three tiers",
    "Feasibility gating precedes any cost comparison of network layouts",
    "Topology flexibility at renewal is worth under 0.5% of life-cycle cost",
    "Residual re-layout value concentrates in larger, sparser clusters",
    "Incumbent-admissible estimand prevents inflated redesign savings",
]
for h in hl:
    assert len(h) <= 85, h
open(f"{ROOT}/manuscript/highlights_ceus.txt", "w").write("\n".join(hl) + "\n")

# ---------------- cover letter ----------------
cl = f"""Cover letter — Computers, Environment and Urban Systems

Dear Editor,

I am pleased to submit "The Value of Topology Flexibility at
Infrastructure Renewal: A Feasibility-Gated Geospatial Assessment of
Low-Voltage Re-Layout During Undergrounding" for consideration as a
research article in Computers, Environment and Urban Systems.

Cities renewing aging overhead lines by undergrounding routinely pay
for the corridor but never price the design option it opens: is
re-drawing the hidden local network worth anything once the trench is
committed? This paper turns that question into a computable geospatial
decision problem and answers it end to end. I reconstruct legacy
low-voltage networks from open municipal and OpenStreetMap data in
three observability tiers, impose engineering feasibility gates before
any cost comparison, validate the linearized electrical model against
AC power flow on all {fmt(mv['ac_n_feeders'],0)} feasible synthesized
feeders, and define an incumbent-admissible estimand — the value of
topology flexibility — that keeps the inherited layout inside the
choice set rather than assuming it away.

The contribution to CEUS's community is, I believe, threefold. First,
a decision layer on top of the geospatial reconstruction literature:
the reconstructed network is treated not as an answer but as the
choice set over which a renewal option is priced. Second, an honest
bounded result: the option value is under 0.5% of life-cycle cost in
all three study zones, and I show that naive incumbent-excluding
contrasts manufacture double-digit artifacts that disappear under
correct estimand design. Third, a spatially structured finding in the
journal's own idiom: residual re-layout value concentrates in larger,
sparser, lower-density feeder clusters — a descriptive screening rule
that tells planners where detailed re-layout engineering is worth
commissioning. The method, data provenance, and full reproducibility
package accompany the submission.

The manuscript is original, is not under consideration elsewhere, and
has been approved for submission.

Sincerely,
Tatsuki Onishi
"""
open(f"{ROOT}/manuscript/cover_letter_ceus.txt", "w").write(cl)

# ---------------- declarations ----------------
decl = """Declarations
- Declaration of generative AI: during the preparation of this work the
  author used an AI coding assistant for code scaffolding and drafting.
  The author reviewed and edited all content and takes full
  responsibility for the publication.
- Conflict of interest: none declared. Funding: none.
- Data availability: see data_code_availability_ceus.txt.
"""
open(f"{ROOT}/manuscript/declarations_ceus.txt", "w").write(decl)

dca = """Data and code availability
All analysis scripts, the source registry, checksums, and
machine-readable result tables are provided in the supplementary
repository snapshot at github.com/bougtoir/underground-topology-segan
(reproduce-only scope). Municipal open data were used under their
stated licenses (CC-BY for Kudamatsu); OpenStreetMap data are used
under ODbL. Processed intermediate tables are included in the
submission package; raw municipal extracts are retained locally per
license terms.
"""
open(f"{ROOT}/manuscript/data_code_availability_ceus.txt", "w").write(dca)

# ---------------- supplement ----------------
sup_doc = Document()
sup_doc.styles["Normal"].font.name = "Times New Roman"
sup_doc.styles["Normal"].font.size = Pt(11)
sup_doc.add_heading("Supplementary material", 0)
sup_doc.add_paragraph(
    "Supplementary Figure S1. NPC of S2 vs the secondary rooftop-PV "
    "overlay S3 by zone (scenario-based upper bound; does not alter the "
    "primary VoF conclusion).")
sup_doc.add_picture(f"{ROOT}/figures/v4_supp_figS1_dg.png", width=Inches(5.5))
sup_doc.add_paragraph(
    "Benchmark note. A second public MV benchmark (Iwamoto 11-bus) "
    "failed to converge in pandapower 3.5.5 and is excluded from "
    "validation statistics; the AC agreement reported in the main text "
    f"covers all {fmt(mv['ac_n_feeders'],0)} feasible synthesized feeders.")
sup_doc.add_paragraph(
    "Additional machine-readable tables (cluster-level results, "
    "morphology descriptors, engineering-assumption sensitivity, "
    "Monte-Carlo sweep) are provided as CSV files in the submission "
    "package.")
sup_doc.save(f"{ROOT}/manuscript/supplement_ceus.docx")

# ---------------- inline docx ----------------
doc = Document(f"{ROOT}/manuscript/manuscript_ceus_v1.docx")
pairs = [
    ("Figure 1 summarizes", ["ceus_fig_pipeline.png",
     "Figure 1. Data-to-decision pipeline: open geospatial inputs, zone construction, feasibility-gated reconstruction, and the incumbent-admissible renewal decision model."]),
    ("(Figure 2)", ["v4_fig1_study_zones.png",
     "Figure 2. Study zones: road graph (grey) and snapped building demand nodes (blue)."]),
    ("Figure 3 shows", ["v4_fig2_topology_example.png",
     "Figure 3. Example feeder cluster: S1 like-for-like replicate vs S2 optimized layout."]),
    ("(Figure 4)", ["v4_fig3_delta.png",
     "Figure 4. Primary estimand VoF by zone vs Monte-Carlo mean of the secondary naive ΔNPC."]),
    ("(Figure 5, Table 3)", ["ceus_fig_morphology.png",
     "Figure 5. Cluster-level VoF vs urban morphology descriptors (descriptive; 94% of clusters have zero VoF)."]),
    ("(Figure 6)", ["ceus_fig_decomposition.png",
     "Figure 6. S1−S2 cost decomposition per zone (accounting decomposition of modeled totals)."]),
]
done = set()
for par in doc.paragraphs:
    for anchor_text, (fn, cap) in pairs:
        if anchor_text in par.text and anchor_text not in done:
            anchor = par._p
            pic = doc.add_paragraph()
            pic.add_run().add_picture(f"{ROOT}/figures/{fn}", width=Inches(5.8))
            capp = doc.add_paragraph(cap)
            anchor.addnext(capp._p); anchor.addnext(pic._p)
            done.add(anchor_text)
print("inserted:", done)
doc.save(f"{ROOT}/manuscript/manuscript_inline_ceus.docx")

# ---------------- readme ----------------
readme = """CEUS SUBMISSION README — portal upload map (Elsevier Editorial Manager)

1. Manuscript (main document): manuscript_ceus_v1.docx
   (inline-figure convenience copy for review: manuscript_inline_ceus.docx)
2. Figures (separate files, both PNG and TIFF supplied):
   ceus_fig_pipeline, v4_fig1_study_zones, v4_fig2_topology_example,
   v4_fig3_delta, ceus_fig_morphology, ceus_fig_decomposition
   → order: pipeline=Fig1, study_zones=Fig2, topology_example=Fig3,
     delta=Fig4, morphology=Fig5, decomposition=Fig6
3. Supplementary material: supplement_ceus.docx (+ v4_supp_figS1_dg)
4. Highlights: highlights_ceus.txt (required by Elsevier)
5. Cover letter: cover_letter_ceus.txt
6. Declarations: declarations_ceus.txt
7. Data/code availability: data_code_availability_ceus.txt
8. Tables: embedded in main docx (Tables 1–4)
9. Fill author name, affiliation, and corresponding-author fields in
   the portal; declare AI use per Elsevier policy (text prepared in
   declarations).
"""
open(f"{ROOT}/manuscript/CEUS_SUBMISSION_README.txt", "w").write(readme)

# ---------------- package: complete manuscript/ subtree ----------------
pkg = f"{ROOT}/CEUS_manuscript_folder_FINAL"
if os.path.exists(pkg):
    shutil.rmtree(pkg)
os.makedirs(pkg)
sub_files = [
    "manuscript_ceus_v1.docx", "manuscript_inline_ceus.docx",
    "supplement_ceus.docx", "cover_letter_ceus.txt", "highlights_ceus.txt",
    "declarations_ceus.txt", "data_code_availability_ceus.txt",
    "CEUS_SUBMISSION_README.txt",
]
os.makedirs(f"{pkg}/manuscript/figures", exist_ok=True)
os.makedirs(f"{pkg}/manuscript/tables", exist_ok=True)
for f in sub_files:
    shutil.copy(f"{ROOT}/manuscript/{f}", f"{pkg}/manuscript/{f}")
figset = ["ceus_fig_pipeline", "v4_fig1_study_zones", "v4_fig2_topology_example",
          "v4_fig3_delta", "ceus_fig_morphology", "ceus_fig_decomposition",
          "v4_supp_figS1_dg"]
for f in figset:
    for ext in ["png", "tiff"]:
        src = f"{ROOT}/figures/{f}.{ext}"
        if os.path.exists(src):
            shutil.copy(src, f"{pkg}/manuscript/figures/{f}.{ext}")
tables = ["analysis/CEUS_MORPHOLOGY_CLUSTER.csv",
          "analysis/CEUS_MORPHOLOGY_CORRELATIONS.csv",
          "analysis/CEUS_BOUNDARY_SENSITIVITY.csv",
          "analysis/manuscript_values_v3.csv", "analysis/scenario_results_v2.csv"]
for f in tables:
    shutil.copy(f"{ROOT}/{f}", f"{pkg}/manuscript/tables/{os.path.basename(f)}")
zpath = f"{ROOT}/CEUS_manuscript_folder_FINAL.zip"
if os.path.exists(zpath):
    os.remove(zpath)
with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
    for dp, _, fs in os.walk(pkg):
        for f in fs:
            fp = os.path.join(dp, f)
            z.write(fp, os.path.relpath(fp, pkg))
print("zip entries:", len(zipfile.ZipFile(zpath).namelist()))
print("done")
