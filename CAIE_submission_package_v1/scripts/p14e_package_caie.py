"""CAIE (Wiley) submission package: cover letter (first-person "I"),
declarations, supplementary docx, inline docx, package dir + zip.
Frozen values from analysis/manuscript_values_v3.csv only."""
import csv, os, zipfile, glob, shutil
from docx import Document
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
mv = {r["key"]: r["value"] for r in
      csv.DictReader(open(f"{ROOT}/analysis/manuscript_values_v3.csv"))}

def fmt(v, nd=2):
    try:
        return f"{float(v):.{nd}f}"
    except Exception:
        return str(v)

# ---------------- cover letter (CAIE; first person; energetic) ----------------
cl = f"""Cover letter — Computer-Aided Civil and Infrastructure Engineering

Dear Editors,

I am pleased to submit "How Much Is Topology Flexibility Worth at
Renewal? A Feasibility-Constrained Life-Cycle Assessment of LV
Re-Layout During Undergrounding" for consideration as an original
research article in Computer-Aided Civil and Infrastructure
Engineering.

Civil infrastructure now faces a large renewal cohort, and every
undergrounding project hides a design question that practice answers by
default: when the trench is open anyway, is it worth re-drawing the
network, or is like-for-like replacement already optimal? This paper
turns that question into a computable decision problem and answers it
end to end. I formulate the renewal window as a choice set in which the
incumbent layout remains an admissible decision, synthesize legacy
low-voltage networks from open municipal data under explicit data
tiers, impose engineering feasibility gates (coincident-load sizing,
ampacity, transformer loading, 6% voltage-drop limit) before any cost
contrast is drawn, and validate the linearized power-flow model against
full AC power flow on all {fmt(mv['ac_n_feeders'],0)} feasible
synthesized feeders.

The central finding is deliberately honest: the value of topology
flexibility was economically equivalent to like-for-like renewal in all
three modeled study zones ({fmt(mv['sano_vof_pct'],2)}%,
{fmt(mv['kudamatsu_vof_pct'],2)}%, {fmt(mv['yasu_vof_pct'],2)}% of
life-cycle cost) under a prespecified 1% equivalence margin. I further
show that a naive forced-redesign contrast — the kind that produces
double-digit headline savings — is an estimand artifact that
disappears once the incumbent is admissible and once engineering
feasibility is gated. For computer-aided infrastructure engineering,
the transferable contribution is a reproducible pipeline that values
re-layout options honestly under open-data constraints: estimand
design, feasibility gating, and reconstruction-ensemble uncertainty,
all released with code and provenance.

I believe this work fits CAIE's scope on computational methods for
infrastructure planning, life-cycle assessment, and decision support
under realistic data constraints.

The manuscript is original, is not under consideration elsewhere, and
has been approved for submission.

Sincerely,
Tatsuki Onishi
"""
open(f"{ROOT}/manuscript/cover_letter_caie.txt", "w").write(cl)

# ---------------- declarations ----------------
decl = """Declarations
- Declaration of generative AI: during the preparation of this work the
  author used an AI coding assistant for code scaffolding and drafting.
  The author reviewed and edited all content and takes full
  responsibility for the publication.
- Data availability: all scripts, the source registry, checksums, and
  machine-readable result tables are provided in the supplementary
  repository snapshot (github.com/bougtoir/underground-topology).
  Municipal open data were used under their stated licenses (CC-BY for
  Kudamatsu).
- Code availability: the full analysis pipeline (scripts/) is included
  in the submission package and the linked public repository.
- Conflict of interest: none declared. Funding: none.
"""
open(f"{ROOT}/manuscript/declarations_caie.txt", "w").write(decl)

# ---------------- supplementary docx ----------------
sup_doc = Document()
sup_doc.styles["Normal"].font.name = "Times New Roman"
sup_doc.styles["Normal"].font.size = Pt(11)
sup_doc.add_heading("Supplementary material", 0)
sup_doc.add_paragraph(
    "Supplementary Figure S1. NPC of S2 vs the secondary rooftop-PV "
    "overlay S3 by zone (scenario-based upper bound; does not alter the "
    "primary VoF conclusion).")
sup_doc.add_picture(f"{ROOT}/figures/v4_supp_figS1_dg.png", width=Inches(5.5))
sup_doc.add_paragraph(
    "Benchmark note. A second public MV benchmark (Iwamoto 11-bus) "
    "failed to converge in pandapower 3.5.5 and is excluded from "
    "validation statistics; the AC agreement reported in the main text "
    f"covers all {fmt(mv['ac_n_feeders'],0)} feasible synthesized feeders.")
sup_doc.add_paragraph(
    "Additional machine-readable tables (cluster-level results, "
    "engineering-assumption sensitivity, Monte-Carlo sweep, "
    "claim-evidence matrix) are provided as CSV files in the submission "
    "package.")
sup_doc.save(f"{ROOT}/manuscript/supplementary_material_caie.docx")

# ---------------- inline docx (figures embedded) ----------------
doc = Document(f"{ROOT}/manuscript/manuscript_caie_v1.docx")
figs = ["v4_fig1_study_zones.png", "v4_fig2_topology_example.png",
        "v4_fig3_delta.png", "v4_fig4_tornado.png"]
caps = {
    "v4_fig1_study_zones.png": "Figure 1. Study zones: road graph (grey) and snapped building demand nodes (blue).",
    "v4_fig2_topology_example.png": "Figure 2. Example feeder cluster: S1 like-for-like replicate vs S2 optimized layout.",
    "v4_fig3_delta.png": "Figure 3. Primary estimand VoF by zone vs Monte-Carlo mean of the secondary naive ΔNPC.",
    "v4_fig4_tornado.png": "Figure 4. Sensitivity to engineering assumptions: secondary naive ΔNPC (bars) vs primary VoF (dots); grey band = prespecified 1% equivalence margin.",
}
insertions = [
    ("(Figures 1–3)", figs[:3]),
    ("(Figure 4)", figs[3:]),
]
done = set()
for par in doc.paragraphs:
    for anchor_text, fl in insertions:
        if anchor_text in par.text and anchor_text not in done:
            anchor = par._p
            for f in fl:
                pic = doc.add_paragraph()
                pic.add_run().add_picture(f"{ROOT}/figures/{f}", width=Inches(5.8))
                cap = doc.add_paragraph(caps[f])
                anchor.addnext(cap._p); anchor.addnext(pic._p)
                anchor = cap._p
            done.add(anchor_text)
print("inserted:", done)
doc.save(f"{ROOT}/manuscript/manuscript_inline_caie.docx")

# ---------------- package ----------------
pkg = f"{ROOT}/CAIE_submission_package_v1"
if os.path.exists(pkg):
    shutil.rmtree(pkg)
os.makedirs(pkg)
files = [
    "manuscript/manuscript_caie_v1.docx", "manuscript/manuscript_inline_caie.docx",
    "manuscript/supplementary_material_caie.docx",
    "manuscript/cover_letter_caie.txt", "manuscript/declarations_caie.txt",
    "analysis/manuscript_values_v3.csv", "analysis/scenario_results_v2.csv",
    "analysis/sensitivity_mc_v2.csv", "analysis/sensitivity_tornado_v2.csv",
    "analysis/dg_results_v2.csv", "analysis/ablation_results.csv",
    "analysis/feasibility_sensitivity.csv", "analysis/original_vs_corrected.csv",
    "analysis/CLAIM_EVIDENCE_MATRIX_FINAL.csv", "analysis/mv_served_nodes.csv",
    "analysis/reconstruction_summary_v2.csv", "analysis/references_verified.csv",
    "analysis/data_audit.csv", "analysis/benchmark_validation.json",
    "analysis/benchmark_ac_solvers.csv", "analysis/benchmark_radial_comparison.csv",
    "analysis/benchmark_iwamoto_diagnosis.csv",
    "analysis/ENGINEERING_ASSUMPTION_SENSITIVITY.csv",
    "analysis/demand_coverage_audit.csv", "analysis/BENEFIT_DECOMPOSITION.csv",
    "analysis/FINAL_V3_VALUE_CHECK.csv",
    "audit_v2/AUDIT_FILE_INVENTORY.csv", "audit_v2/MANUSCRIPT_VALUE_AUDIT.csv",
    "audit_v2/VOLTAGE_ROOT_CAUSE_AUDIT.md", "audit_v2/BENCHMARK_VALIDATION_REPORT.md",
    "audit_v2/KUDAMATSU_MECHANISM_AUDIT.md", "audit_v2/COST_MECHANISM_AUDIT.md",
    "audit_v2/COUNTERFACTUAL_OPTIMIZATION_AUDIT.md", "audit_v2/TERMINOLOGY_AUDIT.md",
    "config/ENGINEERING_FEASIBILITY_SPEC.yaml", "config/ECONOMIC_EQUIVALENCE_SPEC.yaml",
    "config/study.yaml", "PRIMARY_ESTIMAND_SPEC.md",
    "provenance/source_registry.csv", "provenance/decision_log.md",
    "FINAL_REPRODUCIBILITY_REPORT.md", "FINAL_QC_REPORT_v4.md",
    "requirements.lock.txt",
]
for c in ["sano", "kudamatsu", "yasu"]:
    files.append(f"municipality_v4/{c}_planning_support.md")
files += [f"figures/{f}" for f in figs + ["v4_supp_figS1_dg.png"]]
files += [f"scripts/{f}" for f in sorted(glob.glob1(f"{ROOT}/scripts", "*.py"))]
files += [os.path.relpath(os.path.join(dp, f), ROOT)
          for dp, _, fs in os.walk(f"{ROOT}/data/processed") for f in fs]
missing = []
for f in files:
    src = f"{ROOT}/{f}"
    if not os.path.exists(src):
        missing.append(f); continue
    d = os.path.join(pkg, f)
    os.makedirs(os.path.dirname(d), exist_ok=True)
    shutil.copy(src, d)
if missing:
    print("MISSING:", missing)
zpath = f"{ROOT}/CAIE_submission_package_v1.zip"
if os.path.exists(zpath):
    os.remove(zpath)
with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
    for dp, _, fs in os.walk(pkg):
        for f in fs:
            fp = os.path.join(dp, f)
            z.write(fp, os.path.relpath(fp, pkg))
print("zip entries:", len(zipfile.ZipFile(zpath).namelist()))
print("done")
